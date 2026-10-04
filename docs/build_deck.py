"""Builds docs/presentation.pptx from the latest pipeline output.

    python main.py run
    python docs/build_deck.py

Edit NAME / COLLEGE below before submitting.
"""
import json
import sys
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from src import config  # noqa: E402
from src.modules import IndexRebalancer  # noqa: E402
from src.riskengine import load_signals  # noqa: E402

NAME = "Madapati Naga Durga Abhinav Sai"
COLLEGE = "VIT-AP University"

DOCS = ROOT / "docs"
NAVY, GREY, ACCENT = RGBColor(0x1C, 0x28, 0x33), RGBColor(0x56, 0x65, 0x73), RGBColor(0x2E, 0x86, 0xC1)


# --- charts from real output ---------------------------------------------

signals = load_signals(config.SIGNALS_JSONL)
stress = json.loads(config.STRESS_JSON.read_text(encoding="utf-8"))
reb = IndexRebalancer.from_universe(config.UNIVERSE_CSV).run(signals)
snap = pd.DataFrame(reb.snapshot()).sort_values("active_weight")

fig, ax = plt.subplots(figsize=(6, 4.2), dpi=160)
ax.barh(snap["ticker"], snap["active_weight"] * 100,
        color=["#1e8449" if v > 0 else "#c0392b" for v in snap["active_weight"]])
ax.set_xlabel("Active weight vs base (pp)")
ax.set_title("Module A: final tilt after 5 days of news flow")
ax.spines[["top", "right"]].set_visible(False)
fig.tight_layout()
fig.savefig(DOCS / "chart_weights.png")
plt.close(fig)

st = pd.DataFrame([{"label": f"{s['timestamp'][5:10]} {s['event_type'][:13]}",
                    "pnl": s["pnl"] / 1e6} for s in stress])
fig, ax = plt.subplots(figsize=(6, 4.2), dpi=160)
ax.barh(st["label"], st["pnl"], color="#c0392b")
ax.invert_yaxis()
ax.set_xlabel("Stress P&L ($m)")
ax.set_title(f"Module B: {len(stress)} auto-triggered stress tests")
ax.spines[["top", "right"]].set_visible(False)
fig.tight_layout()
fig.savefig(DOCS / "chart_stress.png")
plt.close(fig)

# --- deck --------------------------------------------------------------------

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, body, size=16, bold=False, color=NAVY):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tb.word_wrap = True
    for i, line in enumerate(body if isinstance(body, list) else [body]):
        p = tb.paragraphs[0] if i == 0 else tb.add_paragraph()
        p.text = line
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
        p.space_after = Pt(8)
    return tb


def slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    text(s, 0.6, 0.35, 12, 0.8, title, 30, True)
    if kicker:
        text(s, 0.6, 1.1, 12, 0.5, kicker, 15, color=GREY)
    return s


n_docs = len({s["doc_id"] for s in signals})
n_hi = sum(s["impact_score"] > 7 for s in signals)
worst = min(stress, key=lambda s: s["pnl"])
turnover = sum(h["turnover"] for h in reb.history)

# 1 - title
s = prs.slides.add_slide(BLANK)
text(s, 0.8, 2.2, 11.5, 1.2, "RiskPulse", 54, True)
text(s, 0.8, 3.3, 11.5, 0.8, "An AI/NLP risk engine that turns news and social chatter into "
     "tradeable and stress-testable risk signals", 20, color=GREY)
text(s, 0.8, 5.2, 11.5, 1.2, [NAME, COLLEGE, "S&P Global & Crisil Campus Hackathon 2026"], 16)

# 2 - problem & approach
s = slide("Problem & approach", "Markets move on text long before it shows up in prices or ratings")
text(s, 0.6, 1.8, 6.0, 5, [
    "The problem",
    "- Thousands of headlines and posts per hour; analysts can't read them all",
    "- Unstructured text can't feed a model, a limit or a dashboard",
    "- Risk teams find out about the shock after the P&L does",
], 16)
text(s, 6.9, 1.8, 6.0, 5, [
    "My approach",
    "- One engine, many sources: normalise every feed into a Document",
    "- Score each document on 3 axes: sentiment (-1..1), event type, impact (1-10)",
    "- Publish signals on a bus (JSONL + REST + pub/sub)",
    "- Two consumers prove the signals are usable: a tactical index "
    "rebalancer and an event-driven stress tester",
], 16)

# 3 - architecture
s = slide("System design", "Ingestion -> NLP engine -> signal bus -> downstream modules -> dashboard")
s.shapes.add_picture(str(DOCS / "architecture.png"), Inches(1.4), Inches(1.6), height=Inches(5.7))

# 4 - implementation
s = slide("Implementation highlights")
text(s, 0.6, 1.4, 6.1, 6, [
    "NLP engine (pure Python, explainable)",
    "- Finance lexicon (~250 terms + phrases): 'crushed it' is positive, "
    "'higher for longer' negative, 'liability' isn't",
    "- Negation window, intensifiers, emoji, cashtag linking",
    "- 8-class event taxonomy with weighted triggers + confidence",
    "- Impact = base severity + tone + magnitude ('7%', '$16bn') + shock "
    "words + breadth + source credibility (follower-weighted)",
    "- Pluggable FinBERT backend for sentiment",
], 15)
text(s, 6.9, 1.4, 6.0, 6, [
    "Downstream modules",
    "- A: w ~ base * exp(k * EWMA sentiment), floors/caps, 10% turnover limit",
    "- B: 20-trade synthetic wholesale book (loans, bonds, swaps, CDS, equity)",
    "- Sensitivity-based revaluation: duration, spread duration, delta, FX, PD x LGD",
    "- Scenario picked by event type, scaled by impact; cooldown de-dupes news + social",
    "",
    "Stack: Python, pandas, FastAPI, Streamlit, Plotly, pytest (24 tests)",
], 15)

# 5 - results
s = slide("Key results", f"{n_docs} documents -> {len(signals)} signals, {n_hi} high-impact, "
          f"{len(stress)} stress tests, {len(reb.history)} rebalances")
s.shapes.add_picture(str(DOCS / "chart_weights.png"), Inches(0.5), Inches(1.7), height=Inches(4.6))
s.shapes.add_picture(str(DOCS / "chart_stress.png"), Inches(6.9), Inches(1.7), height=Inches(4.6))
text(s, 0.6, 6.45, 12.2, 1, f"Worst case: '{worst['headline'][:60]}' -> {worst['pnl'] / 1e6:+.1f}m "
     f"({worst['pnl_pct']:+.1%}) on a ${worst['value_before'] / 1e6:.0f}m book. "
     f"Index turnover {turnover:.0%} over the window vs 0% for a static index.", 13, color=GREY)

# 6 - domain impact
s = slide("Domain impact", "Why this matters for a bank or an index provider")
text(s, 0.6, 1.8, 12, 5, [
    "- Early warning: a credit event in the news triggers a portfolio revaluation in seconds, "
    "not at the next daily risk run",
    "- Explainability: every score comes with the words that drove it - auditable for model risk",
    "- One signal, many consumers: the same feed powers trading tilts, limit monitoring and "
    "scenario analysis",
    "- Analyst leverage: a desk watching 18 names and 60 stories sees one ranked feed",
    "- Scales out: swap the CSV adapters for Kafka / vendor APIs without touching the engine",
], 17)

# 7 - limitations
s = slide("Limitations & next steps")
text(s, 0.6, 1.5, 6, 5.5, [
    "Assumptions / gaps",
    "- Lexicon tuned on a small synthetic sample",
    "- First-order (linear) revaluation; no convexity or optionality",
    "- Scenario library is hand-calibrated, not historical",
    "- Entity linking limited to an alias table",
], 16)
text(s, 6.9, 1.5, 6, 5.5, [
    "Next steps",
    "- Fine-tune FinBERT on labelled headlines; ensemble with lexicon",
    "- Calibrate impact against realised abnormal returns",
    "- Historical scenario library (2008, 2020, 2022 rates shock)",
    "- Streaming ingestion (Kafka) + alerting",
    "- Backtest rebalancer vs equal-weight benchmark",
], 16)

out = DOCS / "presentation.pptx"
prs.save(out)
print(f"wrote {out}")

pdf_out = DOCS / "presentation.pdf"
try:
    import win32com.client
    ppt = win32com.client.Dispatch("PowerPoint.Application")
    pres = ppt.Presentations.Open(str(out.resolve()), WithWindow=False)
    pres.SaveAs(str(pdf_out.resolve()), 32)
    pres.Close()
    ppt.Quit()
    print(f"wrote {pdf_out}")
except Exception as e:
    print(f"Note: Could not export PDF via PowerPoint: {e}")
